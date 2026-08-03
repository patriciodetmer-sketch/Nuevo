"""Genera el archivo de merge para una categoría MINVU de entidades patrocinantes de la RM.

Uso:  python3 fuentes/build_merge.py PRIMERA
      python3 fuentes/build_merge.py SEGUNDA
Salida: merge_ds49_categoria_<categoria>.csv y .xlsx
"""
import csv, sys, collections, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CATEGORIA = (sys.argv[1] if len(sys.argv) > 1 else "PRIMERA").upper()
SRC = "inmobiliarias-santiago/entidades_patrocinantes_rm.csv"
slug = "".join(c for c in unicodedata.normalize("NFKD", CATEGORIA.lower()) if not unicodedata.combining(c))
CSV_OUT = f"inmobiliarias-santiago/merge_ds49_categoria_{slug}.csv"
XLSX_OUT = f"inmobiliarias-santiago/merge_ds49_categoria_{slug}.xlsx"

CALLE = {"CALLE","AV","AV.","AVDA","AVDA.","AVENIDA","PASAJE","PSJE","PSJE.","CAMINO","DIAGONAL",
         "GENERAL","GRAL","GRAL.","LOS","LAS","EL","LA","SAN","SANTA","PRESIDENTE","PDTE","PDTE.",
         "PADRE","DR","DR.","DOCTOR","NUEVA","NUEVO","RUTA","AUTOPISTA","ALAMEDA","COSTANERA",
         "MANUEL","JOSE","JOSÉ","VICUÑA","ISIDORA","ROSARIO","CERRO","PARQUE","PLAZA","MONSEÑOR",
         "ALONSO","ANDRES","ANDRÉS","APOQUINDO","AMERICO","AMÉRICO","LIBERTADOR","BERNARDO"}

def partir(blob):
    """Separa nombre de contacto y dirección (heurística: en el PDF vienen fusionados)."""
    tok = blob.split()
    if not tok:
        return "", ""
    for i, t in enumerate(tok[:7]):
        if i >= 1 and t.upper().strip(".,") in CALLE:
            return " ".join(tok[:i]), " ".join(tok[i:])
    corte = min(3, len(tok))
    return " ".join(tok[:corte]), " ".join(tok[corte:])

filas = [r for r in csv.DictReader(open(SRC, encoding="utf-8")) if r["categoria"] == CATEGORIA]
if not filas:
    sys.exit(f"No hay entidades con categoría {CATEGORIA}")

por_correo = collections.defaultdict(list)
for r in filas:
    por_correo[r["correo_principal"].lower()].append(r)

out = []
for correo, grupo in por_correo.items():
    r = grupo[0]
    nombre, direccion = partir(r["contacto_y_direccion"])
    tels = r["telefonos"].split(" / ")
    corta = r["nombre_corto"].strip() or r["razon_social"].strip()
    out.append({
        "correo": correo,
        "entidad": r["razon_social"].strip(),
        "entidad_corta": corta,
        "saludo": f"Estimados {corta}",
        "nombre_contacto_aprox": nombre,
        "rut": r["rut"],
        "direccion_aprox": direccion,
        "comuna": r["comuna"],
        "telefono_1": tels[0] if tels else "",
        "telefonos_todos": r["telefonos"],
        "correos_adicionales": r["correos_adicionales"],
        "categoria": r["categoria"],
        "nota": ("Este correo lo comparten también: " + "; ".join(x["razon_social"] for x in grupo[1:])) if len(grupo) > 1 else "",
    })
out.sort(key=lambda r: r["entidad"])

KEYS = list(out[0].keys())
with open(CSV_OUT, "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=KEYS); w.writeheader(); w.writerows(out)

# ---- xlsx ----
HEAD = ["Correo","Entidad","Entidad corta","Saludo","Contacto (aprox.)","RUT","Dirección (aprox.)",
        "Comuna","Teléfono","Todos los teléfonos","Correos adicionales","Categoría","Nota"]
W = [38,46,34,34,28,15,44,18,18,30,34,12,46]
AZUL, GRIS = "1F3864", "F2F2F2"
thin = Side(style="thin", color="D9D9D9"); border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook(); ws = wb.active; ws.title = "Merge"
ws["A1"] = f"Archivo de merge — DS49, entidades patrocinantes categoría {CATEGORIA.title()} (RM)"
ws["A1"].font = Font(name="Arial", size=14, bold=True, color=AZUL)
ws["A2"] = f"{len(out)} destinatarios únicos · nómina MINVU junio 2026 · una fila = un envío"
ws["A2"].font = Font(name="Arial", size=9, italic=True, color="595959")
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEAD))
H = 4
for c, h in enumerate(HEAD, 1):
    cell = ws.cell(row=H, column=c, value=h)
    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=AZUL)
    cell.alignment = Alignment(vertical="center", wrap_text=True); cell.border = border
ws.row_dimensions[H].height = 30
for i, row in enumerate(out):
    r = H + 1 + i
    for c, k in enumerate(KEYS, 1):
        cell = ws.cell(row=r, column=c, value=row[k])
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (2,7,11,13)))
        cell.border = border
        if i % 2 == 1: cell.fill = PatternFill("solid", fgColor=GRIS)
    mc = ws.cell(row=r, column=1); mc.hyperlink = f"mailto:{row['correo']}"
    mc.font = Font(name="Arial", size=10, color="0563C1", underline="single")
    if row["nota"]:
        ws.cell(row=r, column=13).fill = PatternFill("solid", fgColor="FFEB9C")
ws.freeze_panes = f"A{H+1}"
ws.auto_filter.ref = f"A{H}:{get_column_letter(len(HEAD))}{H+len(out)}"
for c, w_ in enumerate(W, 1): ws.column_dimensions[get_column_letter(c)].width = w_

ins = wb.create_sheet("Instrucciones")
ins["A1"] = "Cómo usar este archivo"
ins["A1"].font = Font(name="Arial", size=14, bold=True, color=AZUL)
comp = sum(1 for r in out if r["nota"])
libres = sum(1 for r in out if r["correo"].split("@")[1] in ("gmail.com","hotmail.com","yahoo.com","outlook.com","yahoo.es"))
txt = [
 "",
 "Qué contiene",
 f"{len(out)} destinatarios únicos: entidades patrocinantes de categoría {CATEGORIA.title()} habilitadas en la Región Metropolitana, según la nómina oficial MINVU de junio 2026. En la nómina son {len(filas)} filas; la diferencia son casillas compartidas por dos entidades, deduplicadas para no enviar dos veces al mismo buzón (marcadas en amarillo en la columna Nota).",
 "",
 "Campos de merge disponibles",
 "{{saludo}} — 'Estimados <nombre corto de la entidad>'. Es el que debes usar para abrir el correo.",
 "{{entidad}} y {{entidad_corta}} — razón social completa y nombre de fantasía.",
 "{{comuna}}, {{telefono_1}}, {{rut}} — para tu registro y el seguimiento telefónico. La comuna es la de la oficina de la entidad, no la del terreno: no la uses en el cuerpo.",
 "",
 "Sobre la columna Contacto (aprox.)",
 "En el PDF de MINVU el nombre de la persona y la dirección vienen en una sola celda, sin separador. La separación es heurística y falla en algunos casos. NO la uses como campo de merge en el saludo. Sirve para saber por quién preguntar cuando llames, confirmándolo en la llamada.",
 "",
 "Antes de enviar",
 "1. Reemplaza los datos del terreno y tu firma en plantilla_merge_ds49.md.",
 "2. Envía desde tu dominio, con SPF y DKIM configurados.",
 "3. Tandas de 20 a 25 por día.",
 "4. Incluye el pie de baja que exige la Ley 19.496 art. 28 B.",
 f"5. {libres} de los {len(out)} correos son de casilla gratuita (Gmail y similares): entidades chicas, donde el teléfono suele rendir más que el correo.",
]
for i, t in enumerate(txt):
    r = 3 + i
    cell = ins.cell(row=r, column=1, value=t)
    bold = t in ("Qué contiene","Campos de merge disponibles","Sobre la columna Contacto (aprox.)","Antes de enviar")
    cell.font = Font(name="Arial", size=10, bold=bold, color=AZUL if bold else "000000")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ins.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    if len(t) > 200: ins.row_dimensions[r].height = 70
    elif len(t) > 100: ins.row_dimensions[r].height = 44
ins.column_dimensions["A"].width = 58
for c in "BCDEFG": ins.column_dimensions[c].width = 14
wb.save(XLSX_OUT)

print(f"categoría {CATEGORIA}: {len(filas)} en la nómina -> {len(out)} destinatarios únicos ({comp} con correo compartido)")
print("comunas:", collections.Counter(r["comuna"] for r in out if r["comuna"]).most_common(6))
print("casilla gratuita:", libres, "| dominio propio:", len(out) - libres)
print("archivos:", CSV_OUT, "|", XLSX_OUT)
